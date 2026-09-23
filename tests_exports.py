import copy, importlib.util, io, json, tempfile, threading, unittest, urllib.request, urllib.error, zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from http.server import ThreadingHTTPServer
import app
from comparison import compare_documents
from export_api import export_document, validate_analysis, FORMATS
from export_xlsx import export_xlsx, MISSING, _chunks

def sample():
    return {"comparison":compare_documents({"before":[{"name":"Редакция 1.md","text":"1. Проверка выполняется.\n2. Удалённая функция."}],
        "after":[{"name":"Редакция 2.md","text":"1. Проверка может выполняться.\n3. Новый контроль."}]})}

class ExportTests(unittest.TestCase):
    def test_xlsx_three_columns_freeze_and_rich_red(self):
        with zipfile.ZipFile(io.BytesIO(export_xlsx(sample()))) as z:
            ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            sheet=ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            self.assertEqual(sheet.find(".//s:pane",ns).attrib["topLeftCell"],"A4")
            self.assertIsNotNone(sheet.find(".//s:autoFilter",ns))
            self.assertTrue(all(len(row.findall("s:c",ns))==3 for row in sheet.findall(".//s:sheetData/s:row",ns)[2:]))
            self.assertIn(MISSING,z.read("xl/worksheets/sheet1.xml").decode())
            self.assertIn('rgb="FFB42335"',z.read("xl/worksheets/sheet1.xml").decode())
            self.assertIsNotNone(sheet.find(".//s:strike",ns))
    def test_xlsx_keeps_very_long_cells_and_linebreaks(self):
        value="Строка документа\n"*2500+"КОНЕЦ"
        rows=[{"before":{"ref":"1","text":value},"after":None,"beforeParts":[{"text":value,"kind":"removed"}],"afterParts":[],"status":"removed","explanation":"Удалено"}]
        a={"comparison":{"documents":{"before":[{"name":"A"}],"after":[{"name":"B"}]},"rows":rows}}
        ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(io.BytesIO(export_xlsx(a))) as z:sheet=ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
        data=sheet.findall(".//s:sheetData/s:row",ns)[3:]
        restored="".join("".join(c.itertext()) for row in data for c in row.findall("s:c",ns) if c.attrib["r"].startswith("A"))
        self.assertEqual(restored,"пункт 1\n"+value)
        self.assertTrue(all(float(row.attrib["ht"])<=409 for row in data))
        self.assertTrue(all(sum(p["text"].count("\n") for p in chunk)<=18 for chunk in _chunks([{"text":value,"kind":"equal"}])))
    def test_xlsx_does_not_create_formulas(self):
        a=sample();a["comparison"]["rows"][0]["explanation"]='=HYPERLINK("http://example.invalid") <script>'
        data=export_xlsx(a)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            text=z.read("xl/worksheets/sheet1.xml").decode()
            self.assertNotIn("<f>",text);self.assertIn("&lt;script&gt;",text)
    def test_rejects_forged_diff_and_format(self):
        a=sample();a["comparison"]["rows"][0]["beforeParts"]=[{"text":"forged","kind":"equal"}]
        with self.assertRaises(ValueError):validate_analysis(a)
        with self.assertRaises(ValueError):export_document(sample(),"../py")
    @unittest.skipUnless(importlib.util.find_spec("docx"),"Install requirements for DOCX")
    def test_docx_is_editable_complete_table_with_redlines(self):
        from docx import Document
        data=export_document(sample(),"docx");doc=Document(io.BytesIO(data))
        self.assertEqual(len(doc.tables[0].columns),3)
        self.assertEqual(len(doc.tables[0].rows),len(sample()["comparison"]["rows"])+1)
        self.assertIn(MISSING,"\n".join(c.text for row in doc.tables[0].rows for c in row.cells))
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml=z.read("word/document.xml").decode()
            for marker in ["w:tblHeader","w:strike","B42335",'w:orient="landscape"']:self.assertIn(marker,xml)
    @unittest.skipUnless(importlib.util.find_spec("reportlab") and importlib.util.find_spec("pypdf"),"Install requirements for PDF")
    def test_pdf_unicode_changes_and_distinct_bold_font(self):
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(export_document(sample(),"pdf")))
        text="".join(p.extract_text() for p in reader.pages)
        self.assertIn("может",text);self.assertIn(MISSING,text);self.assertIn("Редакция 1",text);self.assertNotIn("Редакция 1.md",text)
        fonts={str(obj.get_object().get("/BaseFont","")) for page in reader.pages for obj in page["/Resources"]["/Font"].values()}
        self.assertTrue(any("Roboto-Bold" in font for font in fonts),fonts)
        self.assertTrue(any("Roboto-Regular" in font for font in fonts),fonts)

class BinaryExportHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=ThreadingHTTPServer(("127.0.0.1",0),app.Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.url="http://127.0.0.1:"+str(cls.server.server_port)
    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join()
    def post(self,payload):
        request=urllib.request.Request(self.url+"/api/export-document",json.dumps(payload,ensure_ascii=False).encode(),{"Content-Type":"application/json"})
        with urllib.request.urlopen(request) as r:return json.load(r)
    def test_download_new_formats_with_correct_mime(self):
        for extension in FORMATS:
            if extension=="docx" and not importlib.util.find_spec("docx"):continue
            if extension=="pdf" and not importlib.util.find_spec("reportlab"):continue
            with self.subTest(extension=extension),tempfile.TemporaryDirectory() as folder,patch.object(app,"ROOT",Path(folder)):
                output=self.post({"extension":extension,"analysis":sample()})
                with urllib.request.urlopen(self.url+output["href"]) as response:
                    self.assertEqual(response.headers["Content-Type"],FORMATS[extension])
                    self.assertIn("attachment",response.headers["Content-Disposition"])
                    data=response.read();self.assertTrue(data.startswith(b"%PDF") if extension=="pdf" else data.startswith(b"PK"))
    def test_rejects_invalid_payload(self):
        with self.assertRaises(urllib.error.HTTPError) as cm:self.post({"extension":"xlsx","analysis":{"comparison":{"rows":[]}}})
        self.assertEqual(cm.exception.code,400)

if __name__=="__main__":unittest.main()
