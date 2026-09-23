import base64
import copy
import importlib.util
import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch
import zipfile

import browser_adapter as browser


class BrowserAdapterTests(unittest.TestCase):
    def test_health_reports_browser_processing_without_llm(self):
        health = browser.dispatch("/api/health")
        self.assertTrue(health["browser"])
        self.assertTrue(health["local"])
        self.assertFalse(health["llmConfigured"])
        json.dumps(health)

    def test_control_demo_is_independent_between_requests(self):
        first = browser.dispatch("/api/demo?case=control")
        first["before"][0]["text"] = "changed"
        self.assertNotEqual(browser.dispatch("/api/demo?case=control")["before"][0]["text"], "changed")

    def test_audit_demo_uses_supplied_documents(self):
        demo = browser.dispatch("/api/demo?case=audit")
        self.assertIn("редакция 8", demo["before"][0]["name"])
        self.assertIn("редакция 9", demo["after"][0]["name"])
        self.assertEqual(demo["before"][0]["text"], (browser.ROOT / "data/audit-before.md").read_text(encoding="utf-8"))
        self.assertEqual(demo["after"][0]["text"], (browser.ROOT / "data/audit-after.md").read_text(encoding="utf-8"))

    def test_analysis_preserves_engine_conclusion_and_forces_llm_off(self):
        payload = browser.dispatch("/api/demo?case=control")
        payload["llm"] = True
        with patch.object(browser.backend, "_ollama_refine", side_effect=AssertionError("LLM called")), \
             patch.object(browser.backend.urllib.request, "urlopen", side_effect=AssertionError("network called")), \
             patch.object(Path, "write_text", side_effect=AssertionError("document write")), \
             patch.object(Path, "write_bytes", side_effect=AssertionError("document write")):
            result = browser.dispatch("/api/analyze", payload)
        self.assertTrue(payload["llm"])
        self.assertFalse(result["engine"]["llmUsed"])
        self.assertGreater(result["comparison"]["stats"]["total"], 0)
        self.assertTrue(result["conclusion"]["introduction"])
        self.assertTrue(result["conclusion"]["result"])
        self.assertTrue(all(finding["evidence"] for finding in result["findings"]))
        self.assertIn("elapsedMs", result)
        json.dumps(result)

    def test_extract_keeps_good_files_and_reports_bad_files(self):
        result = browser.dispatch("/api/extract", {"files": [
            {"name": "Редакция.txt", "content": base64.b64encode("1.1. Отдел ведёт реестр договоров.".encode()).decode()},
            {"name": "broken.docx", "content": "not base64!"},
            {"name": "old.doc", "content": base64.b64encode(b"old document").decode()},
        ]})
        self.assertEqual(result["files"][0]["name"], "Редакция")
        self.assertIn("реестр договоров", result["files"][0]["text"])
        self.assertTrue(result["files"][1]["error"])
        self.assertTrue(result["files"][2]["error"])

    def test_limits_reject_empty_or_oversized_requests(self):
        for payload in ({"files": []}, {"files": [{}] * 31}, {"files": [None]}):
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                browser.dispatch("/api/extract", payload)
        sample = browser.dispatch("/api/demo?case=control")
        for documents in ([], [{}], [{"text": "  "}], [{"text": "x"}] * 31):
            with self.subTest(documents=documents), self.assertRaises(ValueError):
                browser.dispatch("/api/analyze", {**sample, "before": documents})
        with patch.object(browser, "MAX_BODY", 10), self.assertRaises(ValueError):
            browser.dispatch("/api/extract", {"files": [{"name": "too big"}]})
        with self.assertRaises(ValueError):
            browser.dispatch("/api/analyze", None)

    def test_invalid_export_format_and_unknown_route_are_rejected(self):
        with self.assertRaises(ValueError):
            browser.dispatch("/api/export-document", {"extension": "../../py", "analysis": {}})
        with self.assertRaises(ValueError):
            browser.dispatch("/api/unknown")


class BrowserExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.analysis = browser.dispatch("/api/analyze", {
            "before": [{"name": "Редакция 1", "text": "1. Отдел проверяет платежи.\n2. Отдел ведёт реестр договоров."}],
            "after": [{"name": "Редакция 2", "text": "1. Отдел может проверять платежи.\n2. Отдел ведёт реестр договоров."}],
        })

    def exported(self, extension):
        result = browser.dispatch("/api/export-document", {"extension": extension, "analysis": self.analysis})
        self.assertEqual(result["encoding"], "base64")
        self.assertEqual(result["extension"], extension)
        self.assertEqual(result["name"], "Versa-сравнение." + extension)
        self.assertEqual(result["mime"], browser.FORMATS[extension])
        json.dumps(result)
        return base64.b64decode(result["content"], validate=True)

    def test_xlsx_export_is_complete_ooxml(self):
        with zipfile.ZipFile(io.BytesIO(self.exported("xlsx"))) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn("может", xml)
        self.assertIn("Редакция 1", xml)
        self.assertIn("B42335", xml)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "DOCX optional library is unavailable")
    def test_docx_export_has_editable_table(self):
        from docx import Document
        doc = Document(io.BytesIO(self.exported("docx")))
        self.assertEqual(len(doc.tables[0].columns), 3)
        self.assertIn("может", " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells))

    @unittest.skipUnless(importlib.util.find_spec("reportlab") and importlib.util.find_spec("pypdf"), "PDF optional libraries are unavailable")
    def test_pdf_export_retains_cyrillic_text(self):
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(self.exported("pdf")))
        self.assertIn("может", " ".join(page.extract_text() for page in pdf.pages))


if __name__ == "__main__":
    unittest.main()
