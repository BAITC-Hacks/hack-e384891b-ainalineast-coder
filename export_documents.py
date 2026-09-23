"""Word and PDF comparative tables with complete source text and marked edits."""
from __future__ import annotations

from html import escape
from io import BytesIO
from pathlib import Path
import re
import threading
from legal_text import display_name, citation

MISSING = "Данный пункт в документе отсутствует."
RED = "B42335"
BLUE = "0D4277"
STATUS = {
    "unchanged": "Положение сохранено без изменений", "changed": "Положение изменено",
    "added": "Включено новое положение", "removed": "Положение исключено",
    "moved": "Изменена нумерация или место положения",
}
_PDF_LOCK = threading.RLock()
_FONTS_READY = False


def _text(value):
    # Office XML and PDF paragraph markup cannot contain these control codes.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "\ufffd", str(value or ""))


def _comparison(analysis):
    comparison = analysis.get("comparison") or {}
    if not isinstance(comparison.get("rows"), list):
        raise ValueError("Нет сравнительной таблицы для экспорта.")
    return comparison


def _names(comparison, side):
    docs = comparison.get("documents", {}).get(side, [])
    return "\n".join(_text(display_name(d.get("name"))) for d in docs if d.get("name")) or (
        "Документ 1" if side == "before" else "Документ 2"
    )


def _parts(row, side):
    clause = row.get(side)
    parts = row.get(side + "Parts")
    if isinstance(parts, list) and "".join(str(p.get("text", "")) for p in parts) == str((clause or {}).get("text", "")):
        return parts
    return [{"text": (clause or {}).get("text", ""), "kind": "equal"}]


def _explanation(row, index):
    status = STATUS.get(row.get("status"), "Сопоставление")
    explanation = _text(row.get("explanation"))
    return f"{index}. {status}" + ("\n" + explanation if explanation else "")


def export_docx(analysis):
    """Return an editable A3 landscape DOCX; rows may span pages without clipping."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    comparison = _comparison(analysis)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(420), Mm(297)
    section.top_margin = section.bottom_margin = Mm(14)
    section.left_margin = section.right_margin = Mm(15)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(4)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(24)
    doc.core_properties.title = "Versa Сравнение документов"
    doc.core_properties.author = "Versa"
    doc.add_paragraph("Versa Сравнение документов", style="Title")
    doc.add_paragraph(
        "Полная сравнительная таблица в последовательности исходных документов. "
        "Красным жирным выделены дополнения и изменения; удалённый текст зачёркнут. "
        f"Строк в таблице: {len(comparison['rows'])}."
    )

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Mm(148.2), Mm(148.2), Mm(93.6)]
    for column, width in zip(table.columns, widths):
        column.width = width
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "bottom", "left", "right"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:w"), "110")
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    properties.append(margins)

    def shade(cell, fill):
        node = OxmlElement("w:shd")
        node.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(node)

    headers = ["Документ 1\n" + _names(comparison, "before"),
               "Документ 2\n" + _names(comparison, "after"), "Содержание изменений и дополнений"]
    for cell, width, title in zip(table.rows[0].cells, widths, headers):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, BLUE)
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(11)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)

    for index, row in enumerate(comparison["rows"], 1):
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if index % 2 == 0:
                shade(cell, "F2F7FC")
        for column, side in enumerate(("before", "after")):
            cell = cells[column]
            clause = row.get(side)
            paragraph = cell.paragraphs[0]
            if clause is None:
                run = paragraph.add_run(MISSING)
                run.italic = True
                run.font.color.rgb = RGBColor.from_string("66758A")
                continue
            if len(comparison.get("documents", {}).get(side, [])) > 1:
                paragraph.add_run(_text(display_name(clause.get("document"))))
                paragraph = cell.add_paragraph()
            ref = _text(citation(clause))
            if ref:
                run = paragraph.add_run(ref)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(BLUE)
                paragraph = cell.add_paragraph()
            for part in _parts(row, side):
                run = paragraph.add_run(_text(part.get("text")))
                if part.get("kind") in ("added", "removed"):
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(RED)
                    run.font.strike = part.get("kind") == "removed"
        paragraph = cells[2].paragraphs[0]
        title, _, detail = _explanation(row, index).partition("\n")
        paragraph.add_run(title).bold = True
        if detail:
            cells[2].add_paragraph(detail)

    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Versa  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _pdf_fonts():
    global _FONTS_READY
    if _FONTS_READY:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    assets = Path(__file__).resolve().parent / "assets"
    pdfmetrics.registerFont(TTFont("VersaText", str(assets / "Roboto-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("VersaText-Bold", str(assets / "Roboto-Bold.ttf")))
    pdfmetrics.registerFontFamily("VersaText", normal="VersaText", bold="VersaText-Bold",
                                italic="VersaText", boldItalic="VersaText-Bold")
    _FONTS_READY = True


def _markup(text):
    return escape(_text(text), quote=False).replace("\n", "<br/>")


def export_pdf(analysis):
    """Return an A3 PDF with embedded Cyrillic fonts and splittable table rows."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    comparison = _comparison(analysis)
    # ReportLab dynamic font state is shared; serialize registration and building.
    with _PDF_LOCK:
        _pdf_fonts()
        stream = BytesIO()
        page_width, page_height = landscape(A3)
        margin = 15 * mm
        width = page_width - 2 * margin - 12
        doc = SimpleDocTemplate(stream, pagesize=(page_width, page_height),
                                leftMargin=margin, rightMargin=margin,
                                topMargin=14 * mm, bottomMargin=14 * mm,
                                title="Versa Сравнение документов", author="Versa")
        body = ParagraphStyle("VersaBody", fontName="VersaText", fontSize=10.5,
                              leading=14, alignment=TA_LEFT, splitLongWords=True,
                              spaceAfter=3, allowWidows=1, allowOrphans=1)
        header = ParagraphStyle("VersaHeader", parent=body, fontName="VersaText-Bold",
                                textColor=colors.white, fontSize=11, leading=14)
        title_style = ParagraphStyle("VersaTitle", parent=body, fontName="VersaText-Bold",
                                     fontSize=23, leading=28, spaceAfter=10)
        story = [Paragraph("Versa Сравнение документов", title_style),
                 Paragraph("Полная сравнительная таблица в последовательности исходных документов. "
                           "Красным жирным выделены дополнения и изменения; удалённый текст зачёркнут. "
                           f"Строк в таблице: {len(comparison['rows'])}.", body), Spacer(1, 10)]
        data = [[Paragraph("Документ 1<br/>" + _markup(_names(comparison, "before")), header),
                 Paragraph("Документ 2<br/>" + _markup(_names(comparison, "after")), header),
                 Paragraph("Содержание изменений и дополнений", header)]]
        for index, row in enumerate(comparison["rows"], 1):
            result = []
            for side in ("before", "after"):
                clause = row.get(side)
                if clause is None:
                    result.append(Paragraph('<font color="#66758A">' + MISSING + "</font>", body))
                    continue
                fragments = []
                for part in _parts(row, side):
                    piece = _markup(part.get("text"))
                    if part.get("kind") in ("added", "removed"):
                        piece = '<font color="#' + RED + '"><b>' + piece + "</b></font>"
                        if part.get("kind") == "removed":
                            piece = "<strike>" + piece + "</strike>"
                    fragments.append(piece)
                ref = _markup(citation(clause))
                text = ('<font color="#' + BLUE + '"><b>' + ref + "</b></font><br/>") if ref else ""
                if len(comparison.get("documents", {}).get(side, [])) > 1:
                    text = _markup(display_name(clause.get("document"))) + "<br/>" + text
                result.append(Paragraph(text + "".join(fragments) or " ", body))
            heading, _, detail = _explanation(row, index).partition("\n")
            result.append(Paragraph("<b>" + _markup(heading) + "</b>" +
                                    ("<br/>" + _markup(detail) if detail else ""), body))
            data.append(result)
        col_widths = [width * .38, width * .38, width * .24]
        header_height = max(cell.wrap(w - 18, 1000000)[1]
                            for cell, w in zip(data[0], col_widths)) + 18
        # Split paragraph content before pagination. ReportLab in-row splitting
        # can mutate repeated headers and clip them on later pages.
        content_height = page_height - 28 * mm - 12 - header_height - 20
        if content_height < 40:
            raise ValueError("Названия документов слишком длинные для заголовка PDF-таблицы.")

        def chunks(paragraph, cell_width):
            result, pending = [], [paragraph]
            while pending:
                current = pending.pop(0)
                if current.wrap(cell_width - 18, content_height)[1] <= content_height:
                    result.append(current)
                    continue
                pieces = current.split(cell_width - 18, content_height)
                if len(pieces) < 2:
                    raise ValueError("Не удалось разместить фрагмент текста в PDF.")
                result.append(pieces[0])
                pending[0:0] = pieces[1:]
            return result

        safe_data = [data[0]]
        for row in data[1:]:
            split_cells = [chunks(cell, w) for cell, w in zip(row, col_widths)]
            for part in range(max(len(c) for c in split_cells)):
                safe_data.append([cell[part] if part < len(cell) else Paragraph(" ", body)
                                  for cell in split_cells])
        table = LongTable(safe_data, colWidths=col_widths, repeatRows=1,
                          splitByRow=1, splitInRow=0, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLUE)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F7FC")]),
            ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
        ]))
        story.append(table)

        def footer(canvas, document):
            canvas.saveState()
            canvas.setFont("VersaText", 8)
            canvas.setFillColor(colors.HexColor("#66758A"))
            canvas.drawRightString(page_width - margin, 7 * mm, f"Versa  |  {document.page}")
            canvas.restoreState()

        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        return stream.getvalue()
